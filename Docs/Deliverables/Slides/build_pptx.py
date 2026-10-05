import os, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image, ImageFont

IMG=os.path.join(os.path.dirname(os.path.abspath(__file__)),'images')+'/'
EQ=os.path.join(os.path.dirname(os.path.abspath(__file__)),'equation_images')+'/'
INK=RGBColor(0x1E,0x29,0x3B); GREEN=RGBColor(0x0F,0x8F,0x85); MINT=RGBColor(0xE6,0xF4,0xF2); DEEP=RGBColor(0x0B,0x12,0x20); BRIGHT=RGBColor(0x2D,0xD4,0xBF); AMBER=RGBColor(0xF5,0x9E,0x0B)
GREY=RGBColor(0x64,0x74,0x8B); WHITE=RGBColor(255,255,255); LINE=RGBColor(0xD5,0xDE,0xE8)
BD=os.path.join(os.path.dirname(os.path.abspath(__file__)),'backdrops')+'/'
FONT='Calibri'; HEAD='Georgia'
try: MET=lambda pt,b=False: ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial Bold.ttf' if b else '/System/Library/Fonts/Supplemental/Arial.ttf', int(pt*10))
except Exception: MET=None
warnings=[]

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
count=[0]

def est_lines(text,w_in,pt,bold=False):
    f=MET(pt,bold); maxw=w_in*72*10  # points*10
    lines=0
    for para in text.split('\n'):
        words=para.split(' '); cur=''; n=1
        for wd in words:
            t=(cur+' '+wd).strip()
            if f.getlength(t)>maxw*0.93 and cur: n+=1; cur=wd
            else: cur=t
        lines+=n
    return lines

def textbox(slide,x,y,w,h,paras,size=18,color=INK,bold=False,align=PP_ALIGN.LEFT,font=FONT,anchor=MSO_ANCHOR.TOP,bullets=False,space=4,italic=False,line=1.0,check=True):
    """paras: list of str or list of runs [(text,{bold,italic,color,size})]"""
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.margin_left=tf.margin_right=Inches(0.04); tf.margin_top=tf.margin_bottom=Inches(0.02); tf.vertical_anchor=anchor
    total=0
    for i,p in enumerate(paras):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        para.alignment=align; para.space_after=Pt(space); para.line_spacing=line
        runs=[(p,{})] if isinstance(p,str) else p
        plain=''
        for t,o in runs:
            r=para.add_run(); r.text=t; plain+=t
            r.font.name=o.get('font',font); r.font.size=Pt(o.get('size',size)); r.font.bold=o.get('bold',bold); r.font.italic=o.get('italic',italic)
            r.font.color.rgb=o.get('color',color)
        if bullets:
            pPr=para._p.get_or_add_pPr(); pPr.set('marL',str(int(Inches(0.28)))); pPr.set('indent',str(-int(Inches(0.28))))
            for tag in ('a:buNone','a:buChar'):
                for e in pPr.findall(qn(tag)): pPr.remove(e)
            bc=etree.SubElement(pPr,qn('a:buClr')); sc=etree.SubElement(bc,qn('a:srgbClr')); sc.set('val','007A3D')
            bu=etree.SubElement(pPr,qn('a:buChar')); bu.set('char','•')
        wid=w-0.1-(0.28 if bullets else 0)
        total+=est_lines(plain,wid,size,bold)*size*1.2*line/72+space/72
    if check and total>h+0.05: warnings.append(f'slide {count[0]}: text overflow {total:.2f}>{h:.2f} :: {str(paras)[:50]}')
    return tb

def rect(slide,x,y,w,h,fill=MINT,line=None,shape=MSO_SHAPE.ROUNDED_RECTANGLE,radius=0.06):
    s=slide.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb=line; s.line.width=Pt(0.75)
    s.shadow.inherit=False
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0]=radius
    return s

def picture(slide,path,x,y,w=None,h=None,border=True):
    im=Image.open(path); ar=im.size[1]/im.size[0]
    if w is not None and h is None: h=w*ar
    elif h is not None and w is None: w=h/ar
    p=slide.shapes.add_picture(path,Inches(x),Inches(y),Inches(w),Inches(h))
    if border: p.line.color.rgb=LINE; p.line.width=Pt(0.75)
    return p,w,h

def new_slide(title=None,bar=True):
    count[0]+=1
    s=prs.slides.add_slide(blank)
    s.shapes.add_picture(BD+'content.jpg',0,0,prs.slide_width,prs.slide_height)
    if title is not None:
        s.shapes.add_picture(BD+'header.jpg',0,0,prs.slide_width,Inches(0.95))
        rect(s,0,0.95,13.333,0.045,fill=BRIGHT,shape=MSO_SHAPE.RECTANGLE)
        rect(s,0.55,0.3,0.1,0.38,fill=AMBER,shape=MSO_SHAPE.RECTANGLE)
        textbox(s,0.8,0.13,11,0.7,[title],size=30,color=WHITE,bold=True,font=HEAD,anchor=MSO_ANCHOR.MIDDLE,check=False)
        textbox(s,11.8,7.08,1.3,0.3,[f'{count[0]} / 10'],size=11,color=GREY,align=PP_ALIGN.RIGHT,check=False)
    return s

def card(slide,x,y,w,h,title,body,size=16,bullets=True):
    rect(slide,x,y,w,h,fill=WHITE,line=LINE)
    rect(slide,x,y,0.09,h,fill=GREEN,shape=MSO_SHAPE.RECTANGLE)
    textbox(slide,x+0.2,y+0.08,w-0.3,0.4,[title],size=size+1,bold=True,color=GREEN,check=False)
    textbox(slide,x+0.2,y+0.5,w-0.3,h-0.58,body,size=size,bullets=bullets,space=3)

def eq(slide,name,x,y,w=None,h=None):
    return picture(slide,EQ+name+'.png',x,y,w,h,border=False)

def eqcard(slide,x,y,w,h,title,names,heights):
    rect(slide,x,y,w,h,fill=WHITE,line=LINE)
    textbox(slide,x+0.15,y+0.06,w-0.3,0.35,[title],size=14,bold=True,color=GREEN,check=False)
    cy=y+0.45
    for n,hh in zip(names,heights):
        im=Image.open(EQ+n+'.png'); ww=hh*im.size[0]/im.size[1]
        if ww>w-0.3: ww=w-0.3; hh=ww*im.size[1]/im.size[0]
        picture(slide,EQ+n+'.png',x+(w-ww)/2,cy,ww,hh,border=False); cy+=hh+0.1

def line_arrow(slide,x1,y1,x2,y2,color=GREEN,width=2):
    c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(width)
    ln=c.line._get_or_add_ln(); te=etree.SubElement(ln,qn('a:tailEnd')); te.set('type','triangle')
    return c

# ---------------- 1 cover
count[0]+=1
s=prs.slides.add_slide(blank)
s.shapes.add_picture(BD+'cover.jpg',0,0,prs.slide_width,prs.slide_height)
LIGHT=RGBColor(0xC7,0xD2,0xE0)
textbox(s,1.0,0.75,10.5,0.4,['KHULNA UNIVERSITY OF ENGINEERING & TECHNOLOGY   ·   CSE 4102'],size=13,color=LIGHT,check=False)
rect(s,1.05,1.25,0.9,0.05,fill=AMBER,shape=MSO_SHAPE.RECTANGLE)
textbox(s,1.0,1.55,11,1.6,['Sentinel'],size=88,bold=True,color=WHITE,font=HEAD,check=False)
textbox(s,1.0,3.2,11,1.2,['A Fully Procedural Game Engine','and Open-World Military Shooter'],size=30,color=RGBColor(0xE6,0xFB,0xF8),space=0,check=False)
textbox(s,1.0,4.55,9,0.9,['Nothing is loaded from a model, image or sound file:','everything is generated from code and a seed.'],size=18,color=BRIGHT,italic=True,space=0,check=False)
textbox(s,1.0,6.45,5.5,0.8,[[('Md. Shifat Hasan',{'bold':True,'color':WHITE,'size':19})],[('Roll 2107067  ·  4th Year 1st Term',{'color':LIGHT,'size':13})]],size=14,space=1,check=False)
textbox(s,6.6,6.45,6.0,0.8,[[('Instructors: Md Tajmilur Rahman, Md. Mubtashim Abrar Nihal',{'size':13,'color':LIGHT})],[('Computer Graphics and Image Processing Laboratory  ·  October 7, 2026',{'size':13,'color':LIGHT})]],size=13,align=PP_ALIGN.RIGHT,space=1,check=False)

# ---------------- 2 outline
s=new_slide('Outline')
items=['Problem, objectives and architecture','Procedural shapes and textures','The rendering pipeline','World, physics and gameplay','Results and verification','Challenge, limits and future work','Video tour of the project']
y=1.45
for i,t in enumerate(items):
    rect(s,0.7,y+0.02,0.5,0.5,fill=GREEN,shape=MSO_SHAPE.OVAL)
    textbox(s,0.7,y+0.02,0.5,0.5,[str(i+1)],size=18,bold=True,color=WHITE,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,check=False)
    textbox(s,1.45,y+0.02,5.9,0.5,[t],size=22,anchor=MSO_ANCHOR.MIDDLE,check=False)
    y+=0.78
p,w,h=picture(s,IMG+'showroom.jpg',7.75,1.7,w=5.0)
textbox(s,7.75,1.7+h+0.1,5.0,0.4,['Every material is a code-generated recipe.'],size=14,color=GREY,italic=True,align=PP_ALIGN.CENTER)

# ---------------- 3 problem + architecture
s=new_slide('1. Problem, objectives and architecture')
textbox(s,0.6,1.2,6.4,0.9,['Can a complete, playable 3D game be made by code alone on a student laptop?'],size=21,bold=True,color=INK)
card(s,0.6,2.2,6.4,1.65,'Two hard limits',['No model, image or precomputed data file','macOS OpenGL 4.1: no compute shaders'],size=18)
textbox(s,0.6,4.0,6.4,0.4,['Objectives'],size=19,bold=True,color=GREEN)
textbox(s,0.6,4.4,6.4,2.6,['Layered engine with downward-only dependencies','All meshes, textures and sounds from code and seeds','Modern renderer: PBR, shadows, sky, post-processing','Playable game: vehicles, ladders, doors, mission','Verified by tests and a 16.7 ms frame budget'],size=17,bullets=True,space=5)
layers=['Game: mission, AI, vehicles, sound','Render: deferred, lights, sky, post','World: collision, terrain','Procedural: noise, textures','GPU: shaders, buffers','Platform: window, input']
y=1.35
for i,t in enumerate(layers):
    r=rect(s,7.5,y,4.7,0.72,fill=RGBColor(0xE4-i*14,0xF6-i*9,0xF3-i*8),line=GREEN)
    textbox(s,7.5,y,4.7,0.72,[t],size=17,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,check=False)
    y+=0.86
line_arrow(s,12.55,1.4,12.55,6.35)
textbox(s,7.5,6.55,5.6,0.4,['A lower layer never knows a higher one.'],size=15,color=GREY,italic=True,align=PP_ALIGN.CENTER)

# ---------------- 4 procedural
s=new_slide('2. Procedural shapes and textures')
p,w,h=picture(s,IMG+'noise_stages.png',2.0,1.15,w=9.3)
textbox(s,2.0,1.15+h+0.03,9.3,0.35,['gradient, fractal, warped, ridged and cellular noise'],size=13,color=GREY,italic=True,align=PP_ALIGN.CENTER)
y0=1.15+h+0.5
eqcard(s,0.6,y0,5.8,6.95-y0,'Fractal noise (tiles exactly)',['fractal'],[1.75]) 
textbox(s,0.75,y0+2.4,5.4,0.35,['Vertex push by a noise deformer'],size=14,bold=True,color=GREEN,check=False)
picture(s,EQ+'push.png',0.85,y0+2.85,w=5.3,border=False)
textbox(s,6.8,y0,6.0,3.4,['18 recipes → 22 materials, baked once at 1024²','Scharr filter turns height into bumps; crevices darken','Objects are tables of primitives plus noise deformers','Soldiers: skeletons with two-bone IK legs'],size=20,bullets=True,space=10)

# ---------------- 5 rendering
s=new_slide('3. The rendering pipeline')
stages=['Sky\ntable','Shadows\n3 cascades','Geometry\nG-buffer','SSAO','Lighting,\nreflect, fog','Smoke,\nfire','Bloom, ACES,\nFXAA']
x=0.6; wbox=1.55; gap=0.2
for i,t in enumerate(stages):
    r=rect(s,x,1.2,wbox,0.8,fill=WHITE,line=GREEN)
    textbox(s,x,1.2,wbox,0.8,t.split('\n'),size=13,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,space=0,check=False)
    if i<len(stages)-1: line_arrow(s,x+wbox,1.6,x+wbox+gap,1.6,width=1.5)
    x+=wbox+gap
textbox(s,0.6,2.2,6.9,0.9,[[('Deferred shading: ',{'bold':True}),('store each pixel\'s colour, normal, shininess and depth first, then light it once; many lights stay cheap.',{})]],size=17)
rect(s,0.6,3.2,6.9,3.7,fill=WHITE,line=LINE)
textbox(s,0.75,3.25,6.5,0.35,['Cook–Torrance reflection'],size=14,bold=True,color=GREEN,check=False)
picture(s,EQ+'ct.png',0.8,3.65,w=6.5,border=False)
picture(s,EQ+'fresnel.png',0.9,4.65,w=3.2,border=False)
textbox(s,0.75,5.15,6.5,0.35,['Fog and tone map'],size=14,bold=True,color=GREEN,check=False)
picture(s,EQ+'fog.png',0.9,5.55,w=3.1,border=False)
picture(s,EQ+'aces.png',4.3,5.45,w=3.0,border=False)
p,w,h=picture(s,IMG+'dusk.jpg',7.85,2.25,w=5.0)
textbox(s,7.85,2.25+h+0.08,5.0,0.4,['Dusk: sun, sky, haze, shadows'],size=14,color=GREY,italic=True,align=PP_ALIGN.CENTER)
textbox(s,7.85,2.25+h+0.6,5.0,1.2,['Sun, moon, lamps, cascaded shadows, SSAO, screen-space reflections, bloom and FXAA'],size=15,color=INK)

# ---------------- 6 world
s=new_slide('4. World, physics and gameplay')
textbox(s,0.6,1.2,6.6,4.1,[[('Sky: ',{'bold':True}),('sun height from astronomy, scattering sky, wind-driven clouds',{})],
 [('Terrain: ',{'bold':True}),('warped ridged noise, streamed in 64 m chunks',{})],
 [('Trees: ',{'bold':True}),('4 species, swaying in the wind',{})],
 [('Vehicles: ',{'bold':True}),('jeep, truck, carrier, tank, helicopter',{})],
 [('Ladders: ',{'bold':True}),('4-state grip machine, hands on rungs',{})],
 [('Mission: ',{'bold':True}),('hack, destroy, rescue, extract; stealth AI; synthesised sound',{})]],size=19,bullets=True,space=7)
rect(s,0.6,5.35,6.6,1.55,fill=WHITE,line=LINE)
picture(s,EQ+'cloud.png',0.85,5.55,w=2.9,border=False)
picture(s,EQ+'bicycle.png',4.9,5.55,h=0.75,border=False)
textbox(s,0.75,6.4,3.2,0.4,['cloud plane and wind'],size=12,color=GREY,italic=True,check=False)
textbox(s,4.4,6.4,2.7,0.4,['bicycle steering'],size=12,color=GREY,italic=True,check=False)
p,w,h=picture(s,IMG+'gate_day.jpg',7.6,1.2,w=5.2)
p,w2,h2=picture(s,IMG+'ladder.jpg',7.6,1.2+h+0.15,w=2.52)
p,w3,h3=picture(s,IMG+'trees.jpg',10.28,1.2+h+0.15,w=2.52)
textbox(s,7.6,1.2+h+0.15+h2+0.05,5.2,0.4,['Day sky, ladder climb, and tree species'],size=13,color=GREY,italic=True,align=PP_ALIGN.CENTER)

# ---------------- 7 results (native chart + table)
s=new_slide('5. Results and verification')
cd=CategoryChartData(); cd.categories=['Sky table','Post','SSAO','Shadows','Lighting','Geometry']
cd.add_series('GPU ms',(0.02,1.21,1.20,1.58,2.79,4.11))
gf=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(0.6),Inches(1.15),Inches(6.4),Inches(3.2),cd); ch=gf.chart
ch.has_legend=False; ch.has_title=True; ch.chart_title.text_frame.text='GPU time per pass (ms)'
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size=Pt(14); ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold=True
pl=ch.plots[0]; pl.gap_width=45; pl.has_data_labels=True; pl.data_labels.font.size=Pt(12); pl.data_labels.number_format='0.00'; pl.data_labels.number_format_is_linked=False
pl.series[0].format.fill.solid(); pl.series[0].format.fill.fore_color.rgb=GREEN
ch.category_axis.tick_labels.font.size=Pt(13); ch.value_axis.tick_labels.font.size=Pt(11); ch.value_axis.has_major_gridlines=False
textbox(s,0.6,4.4,6.4,0.45,[[('10.9 ms',{'bold':True,'color':GREEN}),(' of the 16.7 ms budget at 1080p',{})]],size=20)
rows=[('Verification','Count'),('Unit tests (13 packages)','281'),('TextureCheck (exact tiling)','8,318'),('RenderCheck','318,567'),('GpuCheck','30'),('Playthroughs, 5/5 objectives','2')]
tbl=s.shapes.add_table(len(rows),2,Inches(0.6),Inches(4.95),Inches(6.4),Inches(1.95)).table
tbl.columns[0].width=Inches(4.7); tbl.columns[1].width=Inches(1.7)
for r,(a,b) in enumerate(rows):
    for c,t in enumerate((a,b)):
        cell=tbl.cell(r,c); cell.text=t; para=cell.text_frame.paragraphs[0]; para.runs[0].font.size=Pt(14); para.runs[0].font.name=FONT
        para.alignment=PP_ALIGN.RIGHT if c==1 else PP_ALIGN.LEFT
        cell.margin_top=cell.margin_bottom=Inches(0.02)
        cell.fill.solid(); cell.fill.fore_color.rgb=GREEN if r==0 else (MINT if r%2==0 else WHITE)
        para.runs[0].font.color.rgb=WHITE if r==0 else INK; para.runs[0].font.bold=(r==0)
    tbl.rows[r].height=Inches(0.325)
p,w,h=picture(s,IMG+'hq.jpg',7.9,1.2,w=4.7)
p,w2,h2=picture(s,IMG+'map.jpg',7.9,1.2+h+0.15,w=4.7)
textbox(s,7.9,1.2+h+0.15+h2+0.05,4.7,0.35,['Headquarters, and the tactical map with objectives'],size=13,color=GREY,italic=True,align=PP_ALIGN.CENTER)

# ---------------- 8 challenge
s=new_slide('6. Challenge solved, limits and future work')
textbox(s,0.6,1.15,12.1,1.0,[[('Bug: a bright vertical line at night. ',{'bold':True,'color':GREEN}),('The fog colour sampled the whole sky, so one horizon star painted a streak on every surface in its column. Fix: fog uses the smooth atmosphere only.',{})]],size=18)
p,w,h=picture(s,IMG+'gate_night_bug.jpg',0.6,2.3,w=3.5)
textbox(s,0.6,2.3+h+0.03,3.5,0.35,['Before (streak left of centre)'],size=13,color=GREY,italic=True,align=PP_ALIGN.CENTER)
p,w,h=picture(s,IMG+'gate_night.jpg',0.6,2.3+h+0.5,w=3.5)
textbox(s,0.6,2.3+2*h+0.53,3.5,0.35,['After'],size=13,color=GREY,italic=True,align=PP_ALIGN.CENTER)
card(s,4.85,2.3,7.9,1.3,'Limitations',['Blocky primitive models; no true global illumination; lumpy foliage; tuned by scripts and screenshots'],size=16)
card(s,4.85,3.75,7.9,1.3,'Future work',['Leaf cards, screen-space GI, point-light shadows, multiplayer, a compute-shader port'],size=16)
card(s,4.85,5.2,7.9,1.3,'Take-away',['Most "art" is a few simple ideas, applied consistently and tested'],size=16)

# ---------------- 9 video placeholder (intentionally empty)
s=new_slide('7. Video tour of the project')

# ---------------- 10 references + thanks
s=new_slide('References')
refs=['[1] K. Perlin, "Improving noise," ACM TOG, 21(3), 2002.','[2] S. Worley, "A cellular texture basis function," SIGGRAPH, 1996.','[3] R. Cook, K. Torrance, "A reflectance model for computer graphics," ACM TOG, 1(1), 1982.','[4] B. Walter et al., "Microfacet models for refraction through rough surfaces," EGSR, 2007.','[5] B. Karis, "Real shading in Unreal Engine 4," SIGGRAPH Course, 2013.','[6] S. Hillaire, "A scalable and production ready sky and atmosphere rendering technique," CGF, 39(4), 2020.','[7] R. Dimitrov, "Cascaded shadow maps," NVIDIA, 2007.','[8] M. Oren, S. Nayar, "Generalization of Lambert\'s reflectance model," SIGGRAPH, 1994.',
'[9] K. Narkowicz, "ACES filmic tone mapping curve," 2016.','[10] J. Jimenez, "Next generation post processing in Call of Duty: Advanced Warfare," SIGGRAPH, 2014.','[11] T. Lottes, "FXAA," NVIDIA, 2009.','[12] M. Segal, K. Akeley, The OpenGL Graphics System, v4.1, Khronos, 2010.','[13] I. Quilez, "Domain warping," iquilezles.org.','[14] S. Hargreaves, M. Harris, "Deferred shading," NVIDIA, GDC 2004.','[15] H. Scharr, "Optimal operators in digital image processing," Ph.D., Heidelberg, 2000.','[16] T. Akenine-Möller et al., Real-Time Rendering, 4th ed., CRC Press, 2018.']
textbox(s,0.6,1.15,6.0,5.0,refs[:8],size=16,space=6)
textbox(s,6.85,1.15,6.0,5.0,refs[8:],size=16,space=6)
s.shapes.add_picture(BD+'band.jpg',0,Inches(6.25),prs.slide_width,Inches(1.25))
rect(s,0.7,6.4,0.08,0.9,fill=AMBER,shape=MSO_SHAPE.RECTANGLE)
textbox(s,0.95,6.4,5,0.9,['Thank you'],size=40,bold=True,color=WHITE,font=HEAD,anchor=MSO_ANCHOR.MIDDLE,check=False)
textbox(s,6.0,6.4,6.9,0.9,['Questions are welcome','Md. Shifat Hasan  ·  Roll 2107067'],size=18,color=RGBColor(0xC7,0xD2,0xE0),align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE,space=2,check=False)

out=os.path.join(os.path.dirname(os.path.abspath(__file__)),'Sentinel_Presentation.pptx')
prs.save(out); print('saved',out,len(prs.slides.__iter__.__self__._sldIdLst),'slides'); print('\n'.join(warnings) or 'no overflow warnings')
